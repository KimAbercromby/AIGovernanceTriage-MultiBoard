#!/usr/bin/env python3
"""Regenerate test/fixtures/suite-v3.9.1-contract.json from the suite workbooks and documents.

Usage: python3 scripts/extract-suite-fixture.py <folder with the v3.9.1 .xlsx/.docx sources>

The fixture records, for every workbook sheet or form the tool's exports
target, the exact header row (file, sheet, row) and the controlled lists
(data validations) for the values the tool writes. Tests compare the tool's
export headers and values against it. Needs openpyxl and python-docx.
"""
import json
import sys
from pathlib import Path

import docx
import openpyxl
from openpyxl.utils import get_column_letter

SRC = Path(sys.argv[1])
SUITE = "v3.9.1 (30 September 2026)"


def header(file, sheet, row):
    ws = openpyxl.load_workbook(SRC / file)[sheet]
    values = [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]
    while values and values[-1] is None:
        values.pop()
    formula_cols = [
        get_column_letter(c) for c in range(1, len(values) + 1)
        if isinstance(ws.cell(row + 1, c).value, str) and ws.cell(row + 1, c).value.startswith("=")
    ]
    return {
        "source": {"file": file, "sheet": sheet, "row": row},
        "headers": values,
        "formulaColumns": formula_cols,
    }


def column_a(file, sheet, rows):
    ws = openpyxl.load_workbook(SRC / file)[sheet]
    return {
        "source": {"file": file, "sheet": sheet, "column": "A", "rows": f"{rows[0]}-{rows[-1]}"},
        "labels": {str(r): ws.cell(r, 1).value for r in rows if ws.cell(r, 1).value is not None},
    }


def lists(file, sheet):
    ws = openpyxl.load_workbook(SRC / file)[sheet]
    wb = ws.parent
    out = {}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        f = dv.formula1
        if f.startswith('"'):
            values = f.strip('"').split(",")
        else:  # range reference such as Lists!$A$5:$A$14
            ref_sheet, ref = f.split("!")
            ref = ref.replace("$", "")
            values = [c.value for row in wb[ref_sheet.strip("'")][ref] for c in row if c.value is not None]
        for rng in str(dv.sqref).split():
            out[rng] = values
    return out


def formulas(file, sheet, cells):
    ws = openpyxl.load_workbook(SRC / file)[sheet]
    return {c: ws[c].value for c in cells}


def list_column(file, sheet, col, first, last):
    ws = openpyxl.load_workbook(SRC / file)[sheet]
    return [ws[f"{col}{r}"].value for r in range(first, last + 1) if ws[f"{col}{r}"].value is not None]


def docx_versions():
    d = docx.Document(SRC / "AIG-GOV-03_AI_Governance_Artefact_Index.docx")
    out = {}
    for row in d.tables[0].rows:
        cells = row.cells
        name = cells[0].text
        if name.startswith("AIG-"):
            out[name.split(" — ")[0]] = cells[-1].text.strip()
    return out


def dec02_labels():
    d = docx.Document(SRC / "AIG-DEC-02_Decision_Ready_Paper.docx")
    t = d.tables
    header = [line.split(":")[0].strip() for cell in t[2].rows[0].cells for line in cell.text.split("\n")]
    uc_table = [c.text for c in t[5].rows[0].cells]
    headline = t[4].rows[0].cells[0].text.split("\n")
    return {
        "source": {"file": "AIG-DEC-02_Decision_Ready_Paper.docx", "section": "The template (tables 3-11)"},
        "headerBlock": header,
        "decisionRequired": t[3].rows[0].cells[0].text.split("\n")[0],
        "triageHeadline": headline[0],
        "triageHeadlineText": "\n".join(headline[1:]),
        "ucTable": uc_table,
        "recommendation": t[6].rows[0].cells[0].text.split("\n")[0],
        "bearing": t[7].rows[0].cells[0].text.split("\n")[0],
        "conditionsProposed": t[8].rows[0].cells[0].text.split("\n")[0],
        "fullEvidence": t[9].rows[0].cells[0].text.split("\n")[0],
        "decision": t[10].rows[0].cells[0].text.split("\n")[0],
        "decisionOptions": t[10].rows[0].cells[0].text.split("\n")[1],
    }


def agt06_fields():
    d = docx.Document(SRC / "AIG-AGT-06_Agentic_Action_Decision_Record.docx")
    for t in d.tables:
        if t.rows[0].cells[0].text == "Field":
            return {
                "source": {"file": "AIG-AGT-06_Agentic_Action_Decision_Record.docx", "section": "Controlled fields"},
                "fields": [r.cells[0].text for r in t.rows[1:]],
            }
    raise SystemExit("AGT-06 Controlled fields table not found")


def ops01_section8():
    """AIG-OPS-01 section 8 Rollback and Contingency: [field label, prompt] per row."""
    d = docx.Document(SRC / "AIG-OPS-01_AI_Deployment_and_Rollout_Plan.docx")
    for t in d.tables:
        if t.rows[0].cells[0].text == "Safe-withdrawal / rollback procedure":
            return {
                "source": {"file": "AIG-OPS-01_AI_Deployment_and_Rollout_Plan.docx", "section": "8. Rollback and Contingency"},
                "rows": [[r.cells[0].text, r.cells[1].text] for r in t.rows],
            }
    raise SystemExit("OPS-01 section 8 table not found")


INV04 = "AIG-INV-04_AI_Register_Proposed.xlsx"
DEC04 = "AIG-DEC-04_Gate_Log_Proposed.xlsx"
ASS01 = "AIG-ASS-01_AGPI_Triage_Tool_Proposed.xlsx"
ASS02 = "AIG-ASS-02_AI_Risk_Assessment_Worksheet_Proposed.xlsx"
AGT04 = "AIG-AGT-04_Agent_Record_ASBOM_Proposed.xlsx"
INV05 = "AIG-INV-05_Capabilities_and_System_Map_Proposed.xlsx"
UCV = "UC_ID_Risk_Decision_Current_View.xlsx"
OPS02 = "AIG-OPS-02_AI_Post_Deployment_Monitoring_and_Review_Log_Proposed.xlsx"

fixture = {
    "suite": SUITE,
    "generatedBy": "scripts/extract-suite-fixture.py",
    "versions": docx_versions(),
    "INV-04": {
        "AI Register": {**header(INV04, "AI Register", 3), "lists": lists(INV04, "AI Register")},
        "Assessment summary": {**header(INV04, "Assessment summary", 3), "lists": lists(INV04, "Assessment summary")},
    },
    "DEC-04": {
        "Gate plan": {**header(DEC04, "Gate plan", 3), "lists": lists(DEC04, "Gate plan")},
        "Gate events": {**header(DEC04, "Gate events", 3), "lists": lists(DEC04, "Gate events")},
        "Conditions": {**header(DEC04, "Conditions", 3), "lists": lists(DEC04, "Conditions")},
    },
    "ASS-01": {
        "AGPI Triage": {
            **column_a(ASS01, "AGPI Triage", list(range(4, 23))),
            "lists": lists(ASS01, "AGPI Triage"),
            "priorityAttention": {
                "source": {"file": ASS01, "sheet": "AGPI Triage", "column": "B", "rows": "34-38"},
                "values": list_column(ASS01, "AGPI Triage", "B", 34, 38),
            },
        },
    },
    "ASS-02": {
        "Triage Import": {
            **column_a(ASS02, "Triage Import", list(range(5, 64))),
            "lists": lists(ASS02, "Triage Import"),
        },
        "Risk Assessment": {
            "source": {"file": ASS02, "sheet": "Risk Assessment"},
            "lists": lists(ASS02, "Risk Assessment"),
            "formulas": formulas(ASS02, "Risk Assessment", ["C43", "C82"]),
        },
    },
    "AGT-04": {
        "Agent Record": {**header(AGT04, "Agent Record", 4), "lists": lists(AGT04, "Agent Record")},
        "Capability Vector": {**header(AGT04, "Capability Vector", 4), "lists": lists(AGT04, "Capability Vector")},
        "Runtime Controls": {**header(AGT04, "Runtime Controls", 4), "controlIds": list_column(AGT04, "Runtime Controls", "C", 5, 14)},
        "Authority & Delegations": header(AGT04, "Authority & Delegations", 4),
    },
    "INV-05": {
        s: {**header(INV05, s, 3), "lists": lists(INV05, s)}
        for s in ["System map", "Use cases", "Capabilities", "Relationships"]
    },
    "UCV": {"UC Risk and Decision View": {**header(UCV, "UC Risk and Decision View", 3), "lists": lists(UCV, "UC Risk and Decision View")}},
    "OPS-02": {"Monitoring Log": {**header(OPS02, "Monitoring Log", 4), "lists": lists(OPS02, "Monitoring Log")}},
    "DEC-02": dec02_labels(),
    "AGT-06": agt06_fields(),
    "OPS-01": ops01_section8(),
}

out = Path(__file__).resolve().parent.parent / "test" / "fixtures" / "suite-v3.9.1-contract.json"
out.write_text(json.dumps(fixture, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
print(f"Wrote {out}")
